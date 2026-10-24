"""Clean <PROJECT> per the APPROVED Cleaning Plan.md and the FINAL Issues Requiring Decisions.md.
Copy to "05 Cleaning and Prep/clean_data.py", fill in the TODO blocks (one per plan item, cite its ID), run it.
Reads Raw Data read-only. Writes <prefix>_Clean_draft.xlsx, verifies, then <prefix>_Clean_final.xlsx + <Table>_final.csv."""
import os, sys, json, shutil, hashlib
import pandas as pd
from openpyxl import load_workbook
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
CFG = json.load(open(os.path.join(ROOT, "02 Setup and Log", "project.json"), encoding="utf-8"))
PREFIX = CFG["file_prefix"]
OUT = os.path.join(ROOT, "05 Cleaning and Prep")
DRAFT, FINAL = os.path.join(OUT, f"{PREFIX}_Clean_draft.xlsx"), os.path.join(OUT, f"{PREFIX}_Clean_final.xlsx")
RAW = os.path.join(ROOT, "Raw Data", "TODO_raw_file.xlsx")          # TODO
MEASURES = ["TODO_Sales", "TODO_Profit"]                            # TODO: key measures for control totals

def sha(p):
    h = hashlib.sha256()
    with open(p, "rb") as f:
        for b in iter(lambda: f.read(1 << 20), b""): h.update(b)
    return h.hexdigest()

rel = os.path.relpath(RAW, ROOT).replace("\\", "/")
assert sha(RAW) == CFG["raw_hashes"][rel]["sha256"], "Raw file changed: stop"
T = pd.read_excel(RAW, sheet_name=None)          # or {"Name": pd.read_csv(RAW)}
T = {k: v.copy() for k, v in T.items()}
before = {k: {"rows": len(v), **{m: float(v[m].sum()) for m in MEASURES if m in v}} for k, v in T.items()}
checks = []
def check(name, ok, detail=""):
    checks.append((name, bool(ok), detail)); print(("PASS " if ok else "FAIL ") + name + (f" [{detail}]" if detail else ""))

# ---------------- Class 1 (formatting) ----------------
# TODO DQ-xx: e.g. T["Orders"]["Discount"] = T["Orders"]["Discount"].round(3)
# check("DQ-xx distinct discounts now 27", T["Orders"]["Discount"].nunique() == 27)

# ---------------- Class 2 (user decisions, exactly as recorded) ----------------
# TODO DQ-yy option B: ...

# ---------------- Class 3 (transformations) ----------------
# TODO DQ-zz: derived fields, type conversions, lookups (use .map on a unique key, never a merge that can duplicate rows)

after = {k: {"rows": len(v), **{m: float(v[m].sum()) for m in MEASURES if m in v}} for k, v in T.items()}
for k in before:
    for m, v in before[k].items():
        print(f"{k}.{m}: before {v:,.2f} after {after[k].get(m, float('nan')):,.2f}")
# TODO: add the plan's expected after-values, e.g. check("Orders rows unchanged", after["Orders"]["rows"] == before["Orders"]["rows"])

# ---------------- write: one Excel Table per sheet, minimal formatting ----------------
with pd.ExcelWriter(DRAFT, engine="openpyxl") as xw:
    for k, v in T.items():
        v.to_excel(xw, sheet_name=k, index=False)
wb = load_workbook(DRAFT)
for ws in wb.worksheets:
    ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"
    t = Table(displayName="tbl" + "".join(c for c in ws.title if c.isalnum()), ref=ref)
    t.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=False)
    ws.add_table(t)
wb.save(DRAFT)
re_read = pd.read_excel(DRAFT, sheet_name=None)
for k, v in T.items():
    check(f"draft {k} rows = {len(v)}", len(re_read[k]) == len(v))
if all(ok for _, ok, _ in checks):
    shutil.copyfile(DRAFT, FINAL)
    for k, v in T.items():
        v.to_csv(os.path.join(OUT, f"{k}_final.csv"), index=False)
    print("WROTE", FINAL)
else:
    sys.exit("Checks failed: final file NOT written")
assert sha(RAW) == CFG["raw_hashes"][rel]["sha256"], "Raw file changed during run"
print(json.dumps({"before": before, "after": after, "checks": [(n, ok) for n, ok, _ in checks]}, indent=1))
