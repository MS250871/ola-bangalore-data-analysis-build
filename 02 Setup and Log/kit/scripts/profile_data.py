"""Profile the structure of every file and sheet in Raw Data (step 03). Facts only, no insights.

Usage: python profile_data.py [files...] [--out "03 Data Inspection/profile.json"] [--root PATH]
       (no files = every file in Raw Data)

Per sheet/table: rows, columns, column names, inferred type, non-null and distinct counts, sample values,
min/max for numbers and dates, likely keys (unique non-null columns and 2-column combinations),
date fields, measures (numeric, not ID-like), dimensions (low-cardinality text).
For Excel workbooks also: hidden sheets, hidden rows/columns, merged cells, header row guess,
formula cells vs typed values, more than one table block per sheet (blank-row gaps), Excel Tables,
number formats that suggest currencies or units.
Prints a readable summary and writes full detail to JSON. Never modifies inputs."""
import os, sys, re, argparse, itertools, warnings
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from dakit import find_root, raw_files, write_json
import pandas as pd

warnings.filterwarnings("ignore")
ap = argparse.ArgumentParser()
ap.add_argument("files", nargs="*")
ap.add_argument("--out")
ap.add_argument("--root")
ap.add_argument("--max-sample", type=int, default=5)
A = ap.parse_args()
root = find_root(A.root)
files = [os.path.abspath(f) for f in A.files] or raw_files(root)
out = A.out or os.path.join(root, "03 Data Inspection", "profile.json")


def infer(s):
    nn = s.dropna()
    if nn.empty:
        return "empty"
    if pd.api.types.is_bool_dtype(s):
        return "boolean"
    if pd.api.types.is_datetime64_any_dtype(s):
        return "date" if (nn.dt.normalize() == nn).all() else "datetime"
    if pd.api.types.is_integer_dtype(s):
        return "integer"
    if pd.api.types.is_float_dtype(s):
        return "integer (stored as float)" if (nn == nn.round()).all() else "decimal"
    types = {type(v).__name__ for v in nn.head(5000)}
    if len(types) > 1:
        return "mixed (" + ", ".join(sorted(types)) + ")"
    sv = nn.astype(str)
    if sv.str.fullmatch(r"\d{4}-\d{2}-\d{2}.*").mean() > 0.95:
        return "date as text"
    if sv.str.fullmatch(r"-?\d+(\.\d+)?").mean() > 0.95:
        return "number as text"
    return "text"


def role(name, s, t, n):
    nm = name.lower()
    nun = s.nunique(dropna=True)
    if t.startswith("date") or t == "datetime":
        return "date"
    idlike = bool(re.search(r"(^|[ _])(id|key|code|no|number)$", nm)) or nm.endswith("id")
    if idlike or (nun == n and n > 50 and t in ("text", "integer")):
        return "identifier"
    if t in ("integer", "decimal", "integer (stored as float)"):
        return "dimension (numeric code)" if nun <= 12 and not re.search(r"amount|sales|price|cost|profit|qty|quantity|value|revenue|discount", nm) else "measure"
    if t == "text":
        return "dimension" if nun <= max(50, n * 0.05) else "attribute (high cardinality text)"
    return "other"


def profile_df(df, name):
    n = len(df)
    cols = []
    for c in df.columns:
        s = df[c]
        t = infer(s)
        d = {"column": str(c), "type": t, "non_null": int(s.notna().sum()), "nulls": int(s.isna().sum()),
             "null_pct": round(100 * s.isna().mean(), 2) if n else 0, "distinct": int(s.nunique(dropna=True)),
             "samples": [str(v) for v in s.dropna().unique()[:A.max_sample]]}
        if t in ("integer", "decimal", "integer (stored as float)"):
            d.update({"min": float(s.min()), "max": float(s.max()), "sum": float(s.sum()), "mean": float(s.mean())})
        if t.startswith("date") and pd.api.types.is_datetime64_any_dtype(s):
            d.update({"min": str(s.min()), "max": str(s.max())})
        d["role"] = role(str(c), s, t, n)
        cols.append(d)
    uniq = [c["column"] for c in cols if c["distinct"] == n and c["nulls"] == 0 and n > 0]
    combos = []
    if not uniq and n:
        cands = [c["column"] for c in cols if c["role"] in ("identifier", "dimension", "date") and c["nulls"] == 0][:12]
        for a, b in itertools.combinations(cands, 2):
            if not df.duplicated([a, b]).any():
                combos.append([a, b])
            if len(combos) >= 5:
                break
    dup_rows = int(df.duplicated().sum())
    return {"table": name, "rows": n, "columns": len(df.columns), "column_names": [str(c) for c in df.columns],
            "fully_duplicate_rows": dup_rows, "likely_keys": uniq, "likely_composite_keys": combos,
            "date_fields": [c["column"] for c in cols if c["role"] == "date"],
            "measures": [c["column"] for c in cols if c["role"] == "measure"],
            "dimensions": [c["column"] for c in cols if c["role"].startswith("dimension")],
            "identifiers": [c["column"] for c in cols if c["role"] == "identifier"], "column_detail": cols}


def excel_layout(path):
    from openpyxl import load_workbook
    res = {}
    wb = load_workbook(path, read_only=False, data_only=False)
    for ws in wb.worksheets:
        info = {"state": ws.sheet_state, "dimension": ws.dimensions, "merged_ranges": [str(r) for r in ws.merged_cells.ranges][:20],
                "merged_count": len(ws.merged_cells.ranges), "excel_tables": list(ws.tables.keys()),
                "hidden_columns": [k for k, d in ws.column_dimensions.items() if d.hidden],
                "hidden_rows": sum(1 for k, d in ws.row_dimensions.items() if d.hidden)}
        formulas, blank_rows, header_guess, fmts = 0, [], None, {}
        last_nonblank = None
        for r_idx, row in enumerate(ws.iter_rows(), start=1):
            vals = [c.value for c in row]
            nonblank = [v for v in vals if v not in (None, "")]
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    formulas += 1
                if c.number_format and c.number_format != "General" and isinstance(c.value, (int, float)):
                    fmts[c.number_format] = fmts.get(c.number_format, 0) + 1
            if header_guess is None and len(nonblank) >= max(2, 0.5 * len(vals)) and all(isinstance(v, str) for v in nonblank):
                header_guess = r_idx
            if not nonblank and last_nonblank is not None:
                blank_rows.append(r_idx)
            if nonblank:
                last_nonblank = r_idx
        gaps = [r for r in blank_rows if last_nonblank and r < last_nonblank]
        info.update({"header_row_guess": header_guess, "formula_cells": formulas,
                     "blank_row_gaps_inside_data": gaps[:20], "possible_multiple_tables": bool(gaps),
                     "number_formats": dict(sorted(fmts.items(), key=lambda x: -x[1])[:8])})
        cur = [f for f in info["number_formats"] if re.search(r"[$€£¥₹]|\[\$", f)]
        info["currency_formats"] = cur
        info["mixed_currencies_possible"] = len({re.sub(r"[^$€£¥₹]", "", f) for f in cur}) > 1
        res[ws.title] = info
    return res


report = {"files": []}
for f in files:
    entry = {"file": os.path.relpath(f, root).replace("\\", "/"), "bytes": os.path.getsize(f), "tables": []}
    ext = os.path.splitext(f)[1].lower()
    if ext in (".xlsx", ".xlsm"):
        try:
            entry["excel_layout"] = excel_layout(f)
        except Exception as e:
            entry["excel_layout_error"] = str(e)
    from dakit import read_tables
    for name, df in read_tables(f).items():
        entry["tables"].append(profile_df(df, name))
    report["files"].append(entry)

write_json(out, report)
for e in report["files"]:
    print(f"FILE {e['file']} ({e['bytes']:,} bytes)")
    for sh, li in e.get("excel_layout", {}).items():
        print(f"  sheet '{sh}': state={li['state']} header_row~{li['header_row_guess']} merged={li['merged_count']} "
              f"formulas={li['formula_cells']} hidden_cols={len(li['hidden_columns'])} hidden_rows={li['hidden_rows']} "
              f"tables={li['excel_tables']} multi_table={li['possible_multiple_tables']} currency_formats={li['currency_formats']}")
    for t in e["tables"]:
        print(f"  TABLE {t['table']}: {t['rows']:,} rows x {t['columns']} cols; dup rows={t['fully_duplicate_rows']}; "
              f"keys={t['likely_keys'] or t['likely_composite_keys'] or 'none found'}")
        print(f"    dates={t['date_fields']} measures={t['measures']}")
        print(f"    dimensions={t['dimensions']}")
        for c in t["column_detail"]:
            extra = f" min={c.get('min')} max={c.get('max')}" if "min" in c else ""
            print(f"    - {c['column']}: {c['type']}, {c['role']}, nulls={c['nulls']}, distinct={c['distinct']}{extra}, e.g. {c['samples'][:3]}")
print("WROTE", os.path.relpath(out, root))
