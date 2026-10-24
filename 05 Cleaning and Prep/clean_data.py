"""Clean <PROJECT> per the APPROVED Cleaning Plan.md and the FINAL Issues Requiring Decisions.md.
Copy to "05 Cleaning and Prep/clean_data.py", fill in the TODO blocks (one per plan item, cite its ID), run it.
Reads Raw Data read-only. Writes <prefix>_Clean_draft.xlsx, verifies, then <prefix>_Clean_final.xlsx + <Table>_final.csv."""
import os, sys, json, shutil, hashlib
import pandas as pd
from openpyxl import load_workbook
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
CFG = json.load(open(os.path.join(ROOT, "02 Setup and Log", "project.json"), encoding="utf-8"))
PREFIX = CFG["file_prefix"]
OUT = os.path.join(ROOT, "05 Cleaning and Prep")
DRAFT, FINAL = os.path.join(OUT, f"{PREFIX}_Clean_draft.xlsx"), os.path.join(OUT, f"{PREFIX}_Clean_final.xlsx")
RAW = os.path.join(ROOT, "Raw Data", "ola.csv")
MEASURES = ["Booking_Value", "Ride_Distance", "V_TAT", "C_TAT", "Driver_Ratings", "Customer_Rating"]
import numpy as np

def sha(p):
    h = hashlib.sha256()
    with open(p, "rb") as f:
        for b in iter(lambda: f.read(1 << 20), b""): h.update(b)
    return h.hexdigest()

rel = os.path.relpath(RAW, ROOT).replace("\\", "/")
assert sha(RAW) == CFG["raw_hashes"][rel]["sha256"], "Raw file changed: stop"
RAWTXT = pd.read_csv(RAW, dtype=str, keep_default_na=False, encoding="utf-8-sig")
T = {"ola": RAWTXT.copy()}
T = {k: v.copy() for k, v in T.items()}
NULL_CELLS = int((RAWTXT == "null").sum().sum())
NULL_BY_COL = {c: int((RAWTXT[c] == "null").sum()) for c in RAWTXT.columns}
_t = T["ola"]
for _c in MEASURES:
    _t[_c] = pd.to_numeric(_t[_c].replace("null", np.nan))
before = {k: {"rows": len(v), **{m: float(v[m].sum()) for m in MEASURES if m in v}} for k, v in T.items()}
checks = []
def check(name, ok, detail=""):
    checks.append((name, bool(ok), detail)); print(("PASS " if ok else "FAIL ") + name + (f" [{detail}]" if detail else ""))

# ---------------- Class 1 (formatting) ----------------
# None in the plan.

t = T["ola"]
# ---------------- Class 2 (user decisions, exactly as recorded) ----------------
# DQ-06 A: Booking_Status kept as 4 values (no change)
check("DQ-06 four status values", sorted(t.Booking_Status.unique()) == ["Canceled by Customer","Canceled by Driver","Driver Not Found","Success"])
# DQ-05 B: Success = completed; Incomplete_Rides stays a separate flag
t["Is_Completed"] = (t.Booking_Status == "Success").astype(int)
check("DQ-05 Is_Completed sum = 63,967", t.Is_Completed.sum() == 63967, str(t.Is_Completed.sum()))
check("DQ-05 Success & Incomplete=Yes = 3,926", ((t.Booking_Status=="Success") & (t.Incomplete_Rides=="Yes")).sum() == 3926)
# DQ-09 B: Booking_Value untouched, no revenue-per-km field
bv_before = before["ola"]["Booking_Value"]
# DQ-10 A: flag same-area trips
t["Same_Area_Trip"] = np.where(t.Pickup_Location == t.Drop_Location, "Yes", "No")
check("DQ-10 same-area Yes = 2,121", (t.Same_Area_Trip == "Yes").sum() == 2121)
# DQ-11 B: keep seconds, add minutes
t["V_TAT_Min"] = (t.V_TAT / 60).round(4)
t["C_TAT_Min"] = (t.C_TAT / 60).round(4)
check("DQ-11 V_TAT_Min non-null 63,967", t.V_TAT_Min.notna().sum() == 63967)
check("DQ-11 V_TAT_Min sum ~182,174.77", abs(t.V_TAT_Min.sum() - 182174.77) < 5, f"{t.V_TAT_Min.sum():,.2f}")

# ---------------- Class 3 (transformations) ----------------
# DQ-01: text "null" -> true blank
for c in t.columns:
    if c not in MEASURES:
        t[c] = t[c].replace("null", np.nan)
check("DQ-01 no text 'null' left", int((t.astype(str) == "null").sum().sum()) == 0)
for c in ["Canceled_Rides_by_Customer","Canceled_Rides_by_Driver","Incomplete_Rides_Reason","Payment_Method","Incomplete_Rides"]:
    check(f"DQ-01 blanks in {c} = raw null count", int(t[c].isna().sum()) == NULL_BY_COL[c], str(int(t[c].isna().sum())))
for c in ["V_TAT","C_TAT","Driver_Ratings","Customer_Rating"]:
    check(f"DQ-01 blanks in {c} = 39,057", int(t[c].isna().sum()) == 39057)
# DQ-03: datetime + derived columns, drop Time
dt = pd.to_datetime(t.Date, errors="coerce")
check("DQ-03 all 103,024 dates valid", dt.notna().all())
check("DQ-03 range", dt.min() == pd.Timestamp("2024-07-01 00:00:00") and dt.max() == pd.Timestamp("2024-07-31 23:58:00"), f"{dt.min()} to {dt.max()}")
check("DQ-03 hour matches raw Time", (dt.dt.strftime("%H:%M:%S") == RAWTXT.Time).all())
t["Date"] = dt
t = t.rename(columns={"Date": "Booking_Datetime"})
t["Booking_Date"] = dt.dt.normalize()
t["Booking_Hour"] = dt.dt.hour
t["Weekday"] = dt.dt.day_name()
t["Week_Start"] = (dt.dt.normalize() - pd.to_timedelta(dt.dt.weekday, unit="D"))
t = t.drop(columns=["Time"])
# DQ-02 / DQ-04 (REL-01): drop empty column and image link
t = t.drop(columns=["Unnamed: 20", "Vehicle Images"], errors="ignore")
check("DQ-02/04 dropped columns absent", not {"Unnamed: 20","Vehicle Images","Time"} & set(t.columns))
# DQ-08: distance 0 -> blank
check("DQ-08 zero distances before = 39,057", (t.Ride_Distance == 0).sum() == 39057)
t.loc[t.Ride_Distance == 0, "Ride_Distance"] = np.nan
check("DQ-08 distance blanks = 39,057, sum unchanged, min 1", t.Ride_Distance.isna().sum() == 39057 and t.Ride_Distance.sum() == 1461903 and t.Ride_Distance.min() == 1)
# DQ-07: Revenue only for completed rides
t["Revenue"] = np.where(t.Is_Completed == 1, t.Booking_Value, np.nan)
check("DQ-07 Revenue sum = 35,080,467 over 63,967 rows", t.Revenue.sum() == 35080467 and t.Revenue.notna().sum() == 63967, f"{t.Revenue.sum():,.0f}")
check("DQ-07 quoted non-completed = 21,454,147 and total ties", t.loc[t.Is_Completed==0,"Booking_Value"].sum() == 21454147 and t.Booking_Value.sum() == 56534614)
check("DQ-09 Booking_Value sum unchanged", t.Booking_Value.sum() == bv_before)

# column order
ORDER = ["Booking_Datetime","Booking_Date","Booking_Hour","Weekday","Week_Start","Booking_ID","Booking_Status","Is_Completed","Customer_ID","Vehicle_Type","Pickup_Location","Drop_Location","Same_Area_Trip","V_TAT","C_TAT","V_TAT_Min","C_TAT_Min","Canceled_Rides_by_Customer","Canceled_Rides_by_Driver","Incomplete_Rides","Incomplete_Rides_Reason","Booking_Value","Revenue","Payment_Method","Ride_Distance","Driver_Ratings","Customer_Rating"]
check("column set = 27 expected", sorted(ORDER) == sorted(t.columns) and len(ORDER) == 27, str(len(t.columns)))
t = t[ORDER]
check("rows unchanged 103,024", len(t) == 103024)
T["ola"] = t

after = {k: {"rows": len(v), **{m: float(v[m].sum()) for m in MEASURES if m in v}} for k, v in T.items()}
for k in before:
    for m, v in before[k].items():
        print(f"{k}.{m}: before {v:,.2f} after {after[k].get(m, float('nan')):,.2f}")
# TODO: add the plan's expected after-values, e.g. check("Orders rows unchanged", after["Orders"]["rows"] == before["Orders"]["rows"])

# ---------------- write: one Excel Table per sheet, minimal formatting ----------------
with pd.ExcelWriter(DRAFT, engine="openpyxl") as xw:
    for k, v in T.items():
        v.to_excel(xw, sheet_name=k, index=False)
wb = load_workbook(DRAFT)
for ws in wb.worksheets:
    ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"
    t = Table(displayName="tbl" + "".join(c for c in ws.title if c.isalnum()), ref=ref)
    t.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=False)
    ws.add_table(t)
wb.save(DRAFT)
re_read = pd.read_excel(DRAFT, sheet_name=None)
for k, v in T.items():
    check(f"draft {k} rows = {len(v)}", len(re_read[k]) == len(v))
if all(ok for _, ok, _ in checks):
    shutil.copyfile(DRAFT, FINAL)
    for k, v in T.items():
        v.to_csv(os.path.join(OUT, f"{k}_final.csv"), index=False)
    print("WROTE", FINAL)
else:
    sys.exit("Checks failed: final file NOT written")
assert sha(RAW) == CFG["raw_hashes"][rel]["sha256"], "Raw file changed during run"
print(json.dumps({"before": before, "after": after, "checks": [(n, ok) for n, ok, _ in checks]}, indent=1))
